# mypy: ignore-errors

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.table import Table

from src.core.application.protocol.ports import FileArtifact, ProtocolExporterPort
from src.core.domain.protocol.entities import MeetingProtocol, TaskItem
from src.core.domain.protocol.placeholders import NO_DATA

TITLE = "Протокол встречи"
SECTION_METADATA = "Метаданные"
SECTION_DISCUSSION = "Обсуждение"
SECTION_DECISIONS = "Решения"
SECTION_TASKS = "Задачи"

_TASKS_HEADERS = ("Задача", "Ответственный", "Срок")
_DEFAULT_FONT_SIZE = Pt(11)
_DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_DOCX_FILENAME = "protocol.docx"


class DocxProtocolExporter(ProtocolExporterPort):
    """Адаптер: MeetingProtocol → файл .docx."""

    def export(self, protocol: MeetingProtocol) -> FileArtifact:
        document = self._build_document(protocol)
        buffer = BytesIO()
        document.save(buffer)
        return FileArtifact(
            body=buffer.getvalue(),
            media_type=_DOCX_MEDIA_TYPE,
            filename=_DOCX_FILENAME,
        )

    def export_to_path(self, protocol: MeetingProtocol, path: Path | str) -> None:
        document = self._build_document(protocol)
        document.save(str(path))

    def _build_document(self, protocol: MeetingProtocol) -> Document:
        document = Document()
        self._configure_styles(document)

        document.add_heading(TITLE, level=1)

        document.add_heading(SECTION_METADATA, level=2)
        for line in (
            f"Дата: {protocol.date}",
            f"Участники: {protocol.participants}",
            f"Тема: {protocol.topic}",
        ):
            document.add_paragraph(line, style="List Bullet")

        document.add_heading(SECTION_DISCUSSION, level=2)
        self._add_multiline_text(document, protocol.discussion)

        document.add_heading(SECTION_DECISIONS, level=2)
        if protocol.decisions:
            for index, decision in enumerate(protocol.decisions, start=1):
                document.add_paragraph(f"{index}. {decision}", style="List Number")
        else:
            document.add_paragraph(NO_DATA)

        document.add_heading(SECTION_TASKS, level=2)
        if protocol.tasks:
            self._add_tasks_table(document, protocol.tasks)
        else:
            document.add_paragraph(NO_DATA)

        return document

    @staticmethod
    def _configure_styles(document: Document) -> None:
        normal = document.styles["Normal"]
        normal.font.size = _DEFAULT_FONT_SIZE
        normal.font.name = "Calibri"

        for style_name in ("Heading 1", "Heading 2"):
            heading = document.styles[style_name]
            heading.font.name = "Calibri"

    @staticmethod
    def _add_multiline_text(document: Document, text: str) -> None:
        paragraphs = text.split("\n\n")
        if len(paragraphs) == 1:
            for line in text.splitlines():
                document.add_paragraph(line)
            return

        for block in paragraphs:
            lines = block.splitlines()
            if not lines:
                continue
            paragraph = document.add_paragraph(lines[0])
            for line in lines[1:]:
                paragraph.add_run("\n")
                paragraph.add_run(line)

    @staticmethod
    def _add_tasks_table(document: Document, tasks: tuple[TaskItem, ...]) -> None:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        for index, header in enumerate(_TASKS_HEADERS):
            header_cells[index].text = header
            for paragraph in header_cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = _DEFAULT_FONT_SIZE

        for task in tasks:
            row_cells = table.add_row().cells
            row_cells[0].text = task.title
            row_cells[1].text = task.assignee
            row_cells[2].text = task.deadline

        DocxProtocolExporter._apply_table_font(table)

    @staticmethod
    def _apply_table_font(table: Table) -> None:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = _DEFAULT_FONT_SIZE
