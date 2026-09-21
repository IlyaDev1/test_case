from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.core.abc import SuccessResult
from src.core.application.protocol import (
    ExportProtocolDTO,
    ExportProtocolToDocxUC,
    FileArtifact,
    LinkArtifact,
    ProtocolExporterPort,
    ProtocolTextParserService,
)
from src.core.domain.protocol import NO_DATA, NOT_SPECIFIED, MeetingProtocol
from src.infra.protocol.docx_exporter import (
    SECTION_DECISIONS,
    SECTION_DISCUSSION,
    SECTION_METADATA,
    SECTION_TASKS,
    TITLE,
    DocxProtocolExporter,
)
from src.presentation.fastapi.app import create_app
from src.presentation.fastapi.protocol.responses import artifact_to_response

ROOT = Path(__file__).resolve().parents[2]
EXAMPLES = ROOT / "examples"
FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


@pytest.fixture
def export_uc() -> ExportProtocolToDocxUC:
    return ExportProtocolToDocxUC(
        parser=ProtocolTextParserService(),
        exporter=DocxProtocolExporter(),
    )


def _load_docx(content: bytes) -> Document:
    return Document(BytesIO(content))


def _heading_texts(document: Document) -> list[tuple[int, str]]:
    headings: list[tuple[int, str]] = []
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name == "Heading 1":
            headings.append((1, paragraph.text))
        elif style_name == "Heading 2":
            headings.append((2, paragraph.text))
    return headings


def _section_paragraphs(document: Document, section_title: str) -> list[str]:
    paragraphs = document.paragraphs
    start_index: int | None = None
    for index, paragraph in enumerate(paragraphs):
        if paragraph.style and paragraph.style.name == "Heading 2":
            if paragraph.text == section_title:
                start_index = index + 1
            elif start_index is not None:
                break
    if start_index is None:
        return []

    texts: list[str] = []
    for paragraph in paragraphs[start_index:]:
        if paragraph.style and paragraph.style.name == "Heading 2":
            break
        if paragraph.text:
            texts.append(paragraph.text)
    return texts


def _export_fixture(export_uc: ExportProtocolToDocxUC, name: str) -> bytes:
    text = (FIXTURES / name).read_text(encoding="utf-8")
    result = export_uc.execute(ExportProtocolDTO(text=text))
    assert isinstance(result, SuccessResult)
    assert isinstance(result.data.artifact, FileArtifact)
    return result.data.artifact.body


def test_docx_structure_full_example(export_uc: ExportProtocolToDocxUC) -> None:
    text = (EXAMPLES / "notes_to_protocol.md").read_text(encoding="utf-8")
    result = export_uc.execute(ExportProtocolDTO(text=text))
    assert isinstance(result, SuccessResult)
    assert isinstance(result.data.artifact, FileArtifact)
    document = _load_docx(result.data.artifact.body)

    headings = _heading_texts(document)
    assert headings[0] == (1, TITLE)
    assert [title for _level, title in headings if _level == 2] == [
        SECTION_METADATA,
        SECTION_DISCUSSION,
        SECTION_DECISIONS,
        SECTION_TASKS,
    ]

    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.columns) == 3
    assert [cell.text for cell in table.rows[0].cells] == [
        "Задача",
        "Ответственный",
        "Срок",
    ]
    assert len(table.rows) == 4


def test_docx_missing_assignee_and_deadline(
    export_uc: ExportProtocolToDocxUC,
) -> None:
    content = _export_fixture(export_uc, "protocol_missing_assignee_deadline.md")
    document = _load_docx(content)

    table = document.tables[0]
    rows = [[cell.text for cell in row.cells] for row in table.rows[1:]]

    assert rows[0] == ["Сделать X", NOT_SPECIFIED, NOT_SPECIFIED]
    assert rows[1] == ["Сделать Y", "Боб", NOT_SPECIFIED]
    assert rows[2] == ["Сделать Z", NOT_SPECIFIED, "10.01"]


def test_docx_empty_sections(export_uc: ExportProtocolToDocxUC) -> None:
    content = _export_fixture(export_uc, "protocol_empty_sections.md")
    document = _load_docx(content)

    assert _section_paragraphs(document, SECTION_DECISIONS) == [NO_DATA]
    assert _section_paragraphs(document, SECTION_TASKS) == [NO_DATA]
    assert document.tables == []


def test_api_export_protocol_smoke() -> None:
    text = (EXAMPLES / "notes_to_protocol.md").read_text(encoding="utf-8")
    client = TestClient(create_app())

    response = client.post("/protocol/export", json={"text": text})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(DOCX_MEDIA_TYPE)
    assert 'filename="protocol.docx"' in response.headers["content-disposition"]

    document = _load_docx(response.content)
    assert _heading_texts(document)[0] == (1, TITLE)


class _LinkExporter(ProtocolExporterPort):
    def export(self, protocol: MeetingProtocol) -> LinkArtifact:
        return LinkArtifact(url="https://docs.example/protocol")


def test_export_returns_file_artifact(export_uc: ExportProtocolToDocxUC) -> None:
    text = (EXAMPLES / "notes_to_protocol.md").read_text(encoding="utf-8")
    result = export_uc.execute(ExportProtocolDTO(text=text))
    assert isinstance(result, SuccessResult)

    artifact = result.data.artifact
    assert isinstance(artifact, FileArtifact)
    assert artifact.filename == "protocol.docx"
    assert artifact.body
    assert artifact.media_type.startswith(DOCX_MEDIA_TYPE)


def test_link_artifact_becomes_http_redirect() -> None:
    text = (FIXTURES / "protocol_empty_sections.md").read_text(encoding="utf-8")
    uc = ExportProtocolToDocxUC(
        parser=ProtocolTextParserService(),
        exporter=_LinkExporter(),
    )
    result = uc.execute(ExportProtocolDTO(text=text))
    assert isinstance(result, SuccessResult)

    response = artifact_to_response(result.data.artifact)
    assert response.status_code == 302
    assert response.headers["location"] == "https://docs.example/protocol"
